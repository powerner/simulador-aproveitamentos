#!/usr/bin/env python3
"""
atualizar.py
Lê dados.xlsx e gera index.html atualizado.
Uso: python atualizar.py
Requisitos: pip install openpyxl
"""

import openpyxl, json, re, sys, os
from collections import defaultdict

# Garante que o script trabalha sempre na sua própria pasta
os.chdir(os.path.dirname(os.path.abspath(__file__)))

XLSX = 'dados.xlsx'
TEMPLATE = 'index.html'
SAIDA = 'index.html'

# ── Lê a planilha ──────────────────────────────────────────────────────────
# Aba "Aproveitamentos Normalizados" — colunas lidas pelo nome do cabeçalho:
#   ID Opção | Componente UFRN | Instituição de Origem | Nº Opção na Célula |
#   Componente Cursado | Requer Conjunto (E) | PARECER
# Aba "Componentes" — Componente UFRN | Tipo de Oferta
# Linhas com o mesmo ID Opção formam um conjunto (E): uma linha por componente cursado.
ABA = 'Aproveitamentos Normalizados'
ABA_COMP = 'Componentes'
COLUNAS = {
    'id_opcao':        'ID Opção',
    'comp_ufrn':       'Componente UFRN',
    'instituicao':     'Instituição de Origem',
    'num_opcao':       'Nº Opção na Célula',
    'componente':      'Componente Cursado',
    'requer_conjunto': 'Requer Conjunto (E)',
    'resultado':       'PARECER',
}
COL_VALIDADO = 'Validado pela área?'   # opcional

def txt(v):
    return str(v).replace('\xa0', ' ').strip() if v is not None else ''

def chave(s):
    return ' '.join(txt(s).lower().split())

def ler_planilha(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    cab = {chave(c): i for i, c in enumerate(rows[0]) if c is not None}
    faltando = [n for n in COLUNAS.values() if chave(n) not in cab]
    if faltando:
        print('Erro: colunas não encontradas na planilha:', ', '.join(faltando)); sys.exit(1)
    col = {k: cab[chave(n)] for k, n in COLUNAS.items()}
    col_val = cab.get(chave(COL_VALIDADO))

    tipos = {}
    if ABA_COMP in wb.sheetnames:
        for r in list(wb[ABA_COMP].iter_rows(values_only=True))[1:]:
            if r and txt(r[0]):
                tipos[txt(r[0])] = txt(r[1] if len(r) > 1 else '').upper()
    else:
        print(f'  AVISO: aba "{ABA_COMP}" não encontrada; Tipo de Oferta ficará em branco.')

    dados = []
    for n, r in enumerate(rows[1:], start=2):
        v = lambda k: r[col[k]] if col[k] < len(r) else None
        if v('id_opcao') in (None, '') or not txt(v('comp_ufrn')) or not txt(v('componente')):
            if any(x not in (None, '') for x in r):
                print(f'  AVISO: linha {n} ignorada (sem ID, componente UFRN ou componente cursado).')
            continue
        comp = txt(v('comp_ufrn'))
        if comp not in tipos:
            print(f'  AVISO: linha {n}: "{comp}" não está na aba {ABA_COMP}.')
        dados.append({
            'id_opcao':        int(v('id_opcao')),
            'comp_ufrn':       comp,
            'tipo_oferta':     tipos.get(comp, ''),
            'instituicao':     txt(v('instituicao')),
            'num_opcao':       int(v('num_opcao')) if v('num_opcao') else 1,
            'componente':      txt(v('componente')),
            'requer_conjunto': txt(v('requer_conjunto')).lower() == 'sim',
            'resultado':       txt(v('resultado')).upper() or 'DEFERIDO',
            'validado':        col_val is not None and col_val < len(r) and txt(r[col_val]).lower() == 'sim',
        })
    return dados

# ── Normaliza para JSON do simulador ──────────────────────────────────────
def normalizar(dados):
    tipos = {}
    for r in dados:
        if r['tipo_oferta'] and r['comp_ufrn'] not in tipos:
            tipos[r['comp_ufrn']] = r['tipo_oferta']

    # Agrupa por (componente UFRN, instituição, ID Opção); cada linha do ID é um componente do conjunto
    opcoes_map = defaultdict(lambda: defaultdict(dict))
    for r in dados:
        c, inst, iop = r['comp_ufrn'], r['instituicao'], r['id_opcao']
        op = opcoes_map[c][inst].get(iop)
        if op is None:
            op = opcoes_map[c][inst][iop] = {
                'id_opcao':        iop,
                'num_opcao':       r['num_opcao'],
                'componentes':     [],
                'requer_conjunto': False,
                'resultado':       r['resultado'],
                'validado':        r['validado'],
            }
        elif op['resultado'] != r['resultado']:
            print(f'  AVISO: ID {iop} tem pareceres diferentes entre as linhas; usado INDEFERIDO.')
            op['resultado'] = 'INDEFERIDO'
        if op['validado'] != r['validado'] and op['componentes']:
            print(f'  AVISO: ID {iop} tem linhas validadas e não validadas; considerado não validado.')
            op['validado'] = False
        op['componentes'].append(r['componente'])
        op['requer_conjunto'] = op['requer_conjunto'] or r['requer_conjunto'] or len(op['componentes']) > 1

    estrutura = {}
    for c in sorted(opcoes_map.keys()):
        estrutura[c] = {'tipo_oferta': tipos.get(c, ''), 'instituicoes': {}}
        for inst in sorted(opcoes_map[c].keys()):
            ops = sorted(opcoes_map[c][inst].values(), key=lambda x: (x['num_opcao'], x['id_opcao']))
            ops = [o for o in ops if o['resultado'] in ('DEFERIDO', 'INDEFERIDO')]
            if ops:
                estrutura[c]['instituicoes'][inst] = ops
    return estrutura

# ── Gera HTML ─────────────────────────────────────────────────────────────
def gerar_html(estrutura, template_path, saida_path):
    with open(template_path, 'r', encoding='utf-8') as f:
        html = f.read()

    dados_json  = json.dumps(estrutura, ensure_ascii=False, separators=(',', ':'))
    comps_lista = sorted(estrutura.keys())
    comps_js    = json.dumps(comps_lista, ensure_ascii=False)
    opcoes_comp = '\n'.join([
        '<option value="' + c.replace('"', '&quot;') + '">' + c + '</option>'
        for c in comps_lista
    ])

    # Substitui var DADOS usando delimitadores exatos (evita regex DOTALL guloso)
    marcador_dados = 'var DADOS = {'
    if marcador_dados in html:
        pos_di = html.index(marcador_dados)
        # Encontra o fechamento correto: navega pelos {} para achar o par
        depth = 0
        pos_df = pos_di + len('var DADOS = ')
        i = pos_df
        while i < len(html):
            if html[i] == '{':
                depth += 1
            elif html[i] == '}':
                depth -= 1
                if depth == 0:
                    pos_df = i + 2  # inclui o };
                    break
            i += 1
        html = html[:pos_di] + 'var DADOS = ' + dados_json + ';' + html[pos_df:]
    else:
        print('AVISO: var DADOS nao encontrado no template.')

    # Substitui var COMPS_LISTA
    marcador_comps = 'var COMPS_LISTA = ['
    if marcador_comps in html:
        pos_ci = html.index(marcador_comps)
        pos_cf = html.index('];', pos_ci) + 2
        html = html[:pos_ci] + 'var COMPS_LISTA = ' + comps_js + ';' + html[pos_cf:]
    else:
        print('AVISO: var COMPS_LISTA nao encontrado no template.')

    # Substitui as options do datalist
    html = re.sub(
        r'(<datalist id="lista-comps-ufrn">).*?(</datalist>)',
        r'\g<1>\n' + opcoes_comp + r'\n\g<2>',
        html, count=1, flags=re.DOTALL
    )

    with open(saida_path, 'w', encoding='utf-8') as f:
        f.write(html)

# ── Main ──────────────────────────────────────────────────────────────────
if __name__ == '__main__':
    if not os.path.exists(XLSX):
        print(f'Erro: {XLSX} não encontrado.')
        sys.exit(1)
    if not os.path.exists(TEMPLATE):
        print(f'Erro: {TEMPLATE} não encontrado.')
        sys.exit(1)

    print(f'Lendo {XLSX}...')
    dados = ler_planilha(XLSX)
    print(f'  {len(dados)} linhas lidas.')

    print('Normalizando...')
    estrutura = normalizar(dados)
    total_opts = sum(
        len(ops)
        for v in estrutura.values()
        for ops in v['instituicoes'].values()
    )
    print(f'  {len(estrutura)} componentes UFRN, {total_opts} opções de equivalência.')

    print(f'Gerando {SAIDA}...')
    gerar_html(estrutura, TEMPLATE, SAIDA)
    print(f'Concluído. Arquivo gerado: {SAIDA}')
    print()
    print('Próximo passo: suba o index.html no GitHub para publicar o site atualizado.')